# -*- coding: utf-8 -*-
"""
Script Generator DOCX: UTS_I1_RTConnect_SRS.docx
Standar IEEE 830 / Nigeria MFL Style
Praktikum Pembangunan Perangkat Lunak (Kelas I1) - Universitas Airlangga 2026
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        for r in h.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(26, 115, 232) # Google Blue
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(32, 33, 36)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(95, 99, 104)
    return h

def add_p(doc, text, bold_prefix="", italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(32, 33, 36)
    r = p.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_callout(doc, title, text, border_color="#1A73E8", bg_color="#F8F9FA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color.replace("#", ""))
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color.replace('#','')}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{title}\n")
    r1.font.name = 'Segoe UI'
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(26, 115, 232)
    
    r2 = p.add_run(text)
    r2.font.name = 'Segoe UI'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(60, 64, 67)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_diagram_image(doc, img_path, caption, width_inch=6.0):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inch))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = 'Segoe UI'
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(95, 99, 104)
    else:
        p = doc.add_paragraph(f"[Gambar tidak ditemukan: {img_path}]")
        p.runs[0].font.color.rgb = RGBColor(220, 38, 38)

def format_table(table, col_widths, col_alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for j, cell in enumerate(row.cells):
            cell.width = col_widths[j]
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if i == 0:
                set_cell_background(cell, "E8F0FE")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.name = 'Segoe UI'
                        r.font.size = Pt(9)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(26, 115, 232)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, "F8F9FA")
                for p in cell.paragraphs:
                    if col_alignments and j < len(col_alignments):
                        p.alignment = col_alignments[j]
                    for r in p.runs:
                        r.font.name = 'Segoe UI'
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(32, 33, 36)

print("Helper functions defined!")
